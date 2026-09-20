# -*- coding: utf-8 -*-
"""Project demo PPT: FAE 知识智能治理平台. Visual, spoken-word friendly."""

from __future__ import annotations

import os
import shutil

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt as PptPt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "docs", "AI应用大赛相关材料")
FIG = os.path.join(OUT_DIR, "_review_figs")
OUT_NAME = "复审_PPT_FAE知识智能治理平台.pptx"

NAVY = PptRGB(0x1F, 0x3A, 0x5F)
NAVY2 = PptRGB(0x2C, 0x49, 0x6E)
ORANGE = PptRGB(0xC4, 0x5C, 0x26)
BG = PptRGB(0xF4, 0xF6, 0xF9)
CARD = PptRGB(0xFF, 0xFF, 0xFF)
MUTED = PptRGB(0x5C, 0x67, 0x74)
LINE = PptRGB(0xD6, 0xDD, 0xE6)
WHITE = PptRGB(0xFF, 0xFF, 0xFF)
SOFT = PptRGB(0xE8, 0xEF, 0xF7)

W = 12192000
H = 6858000
HEADER_H = 720000
FOOT_H = 280000


def _set_run(run, *, size=14, bold=False, color=NAVY, name="Microsoft YaHei"):
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _add_text(slide, l, t, w, h, text, *, size=14, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = PptPt(4)
        run = p.add_run()
        run.text = line
        _set_run(run, size=size, bold=bold, color=color)
    return box


def _rect(slide, l, t, w, h, fill, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    if radius:
        sh.adjustments[0] = radius
    return sh


def _accent_bar(slide, l, t, w, h, fill=ORANGE):
    return _rect(slide, l, t, w, h, fill)


def _card(slide, l, t, w, h):
    sh = _rect(slide, l, t, w, h, CARD, radius=0.08)
    sh.line.fill.solid()
    sh.line.color.rgb = LINE
    sh.line.width = Emu(12700)
    return sh


def _header(slide, title, page, total):
    _rect(slide, 0, 0, W, H, BG)
    _rect(slide, 0, 0, W, HEADER_H, NAVY)
    _rect(slide, 0, HEADER_H - 40000, W, 40000, ORANGE)
    _add_text(slide, 360000, 160000, 9000000, 420000, title, size=18, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    _add_text(
        slide,
        9600000,
        180000,
        2200000,
        380000,
        f"{page}  /  {total}",
        size=12,
        color=PptRGB(0xC5, 0xD0, 0xDC),
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _add_text(
        slide,
        360000,
        H - FOOT_H,
        8000000,
        220000,
        "FAE 知识智能治理平台  ·  分类小能手",
        size=10,
        color=MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def _add_pic(slide, path, l, t, max_w, max_h):
    im = Image.open(path)
    iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    x = l + (max_w - w) // 2
    y = t + (max_h - h) // 2
    slide.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))


def _blank():
    prs = Presentation()
    prs.slide_width = Emu(W)
    prs.slide_height = Emu(H)
    return prs


def _slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build() -> str:
    prs = _blank()
    total = 11
    page = 0

    def nxt(title=None):
        nonlocal page
        page += 1
        s = _slide(prs)
        if title:
            _header(s, title, page, total)
        else:
            _rect(s, 0, 0, W, H, BG)
        return s

    # 1 cover
    s = nxt(None)
    _rect(s, 0, 0, 5200000, H, NAVY)
    _rect(s, 0, 0, 90000, H, ORANGE)
    _add_text(s, 420000, 900000, 4500000, 400000, "移远通信  ·  AI 实战应用", size=14, color=PptRGB(0xC5, 0xD0, 0xDC))
    _add_text(s, 420000, 1500000, 4500000, 1400000, "FAE 知识\n智能治理平台", size=36, bold=True, color=WHITE)
    _add_text(
        s,
        420000,
        3200000,
        4500000,
        900000,
        "把按人堆放的 Sharing，治成可检索、可复用的技术资产",
        size=16,
        color=PptRGB(0xD7, 0xE0, 0xEA),
    )
    _add_text(
        s,
        420000,
        5600000,
        4500000,
        600000,
        "团队：分类小能手　　部门：技术支持部\n赛道：运营提效与职能创新",
        size=13,
        color=PptRGB(0xC5, 0xD0, 0xDC),
    )
    metrics = [
        ("17,105", "篇已分类进枢纽"),
        ("17,092", "篇标题标准化"),
        ("169,010", "张附件图已重绑"),
        ("约 1.7", "FTE / 年持续提效"),
    ]
    for i, (n, lab) in enumerate(metrics):
        y = 700000 + i * 1400000
        _card(s, 5600000, y, 6000000, 1220000)
        _accent_bar(s, 5600000, y, 70000, 1220000)
        _add_text(s, 5900000, y + 220000, 5400000, 500000, n, size=28, bold=True, color=ORANGE)
        _add_text(s, 5900000, y + 720000, 5400000, 350000, lab, size=14, color=MUTED)

    # 2 agenda
    s = nxt("今天演示什么")
    items = [
        ("01", "问题", "FAE Sharing 存得多，但找不到、附件进不了检索"),
        ("02", "做法", "飞书做存储，自研 Agent 做抽取、分类、标题、编排、修复"),
        ("03", "效果", "1.7 万篇进枢纽，16.9 万张图可读，检索从翻人名变成找型号"),
        ("04", "下一步", "质量评分 → 检索问答 / 工单辅助；流水线可迁到其他知识空间"),
    ]
    for i, (num, title, body) in enumerate(items):
        y = 1000000 + i * 1300000
        _card(s, 480000, y, 11200000, 1150000)
        _rect(s, 480000, y, 1400000, 1150000, NAVY if i < 3 else ORANGE, radius=0.08)
        _add_text(s, 480000, y + 320000, 1400000, 500000, num, size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text(s, 2100000, y + 220000, 9000000, 400000, title, size=22, bold=True)
        _add_text(s, 2100000, y + 620000, 9000000, 400000, body, size=15, color=MUTED)

    # 3 pain
    s = nxt("现场痛点：存得多 ≠ 用得上")
    pains = [
        ("按人堆放", "20,000+ 项历史 Sharing\n目录按区域/人员沉淀，跨产品线找不到同类案例"),
        ("附件沉睡", "有效步骤在 PDF / Word / PPT 里\n只读飞书正文会漏分类，也搜不到 AT 命令和报错码"),
        ("标题不可读", "日期、流水号、临时描述\n命中后无法判断主题、型号、作者"),
        ("飞书不解治理", "Wiki 提供编辑、目录、基础检索\n不解析附件、不按模组规范分类、不做增量编排"),
    ]
    for i, (t, b) in enumerate(pains):
        col, row = i % 2, i // 2
        x, y = 480000 + col * 5700000, 1050000 + row * 2600000
        _card(s, x, y, 5400000, 2350000)
        _accent_bar(s, x, y, 5400000, 90000)
        _add_text(s, x + 280000, y + 280000, 4900000, 500000, t, size=22, bold=True, color=ORANGE)
        _add_text(s, x + 280000, y + 900000, 4900000, 1200000, b, size=15, color=MUTED)

    # 4 architecture
    s = nxt("方案分层：飞书存储  ·  自研 Agent 治理")
    _add_pic(s, os.path.join(FIG, "fig4_agent_mindmap.png"), 200000, 900000, 11800000, 5400000)

    # 5 agents
    s = nxt("已上线的原生 Agent")
    _add_pic(s, os.path.join(FIG, "fig5_agent_stack.png"), 200000, 880000, 11800000, 5000000)
    _add_text(
        s,
        480000,
        5900000,
        11200000,
        400000,
        "模型：deepseek-v4-flash（经公司网关）　　分类受 QT-SOP 正则与标签树约束，不能自由编造路径",
        size=13,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # 6 extract
    s = nxt("附件如何进入知识库：以 PDF 为例")
    _add_pic(s, os.path.join(FIG, "fig6_pdf_extract.png"), 150000, 860000, 11900000, 5450000)

    # 7 storage recall
    s = nxt("知识如何存储，检索如何变准")
    _add_pic(s, os.path.join(FIG, "fig7_storage_recall.png"), 150000, 860000, 11900000, 5450000)

    # 8 effect
    s = nxt("落地效果（可核验）")
    _add_pic(s, os.path.join(FIG, "fig11_effect.png"), 200000, 850000, 11800000, 3600000)
    deltas = [
        ("归档", "5 → 0.5 分钟/篇", "90%"),
        ("找同类案例", "15 → 3 分钟/次", "80%"),
        ("改标题", "2 → 0 分钟/篇", "100%"),
        ("打开附件判断", "8 → 2 分钟/次", "75%"),
    ]
    for i, (k, v, r) in enumerate(deltas):
        x = 360000 + i * 2920000
        _card(s, x, 4550000, 2750000, 1450000)
        _add_text(s, x + 140000, 4680000, 2470000, 320000, k, size=13, color=MUTED)
        _add_text(s, x + 140000, 5000000, 2470000, 400000, v, size=16, bold=True)
        _add_text(s, x + 140000, 5450000, 2470000, 350000, f"提效 {r}", size=14, bold=True, color=ORANGE)

    # 9 usage
    s = nxt("使用变化：检索、阅读、写作、运营")
    uses = [
        ("一线检索", "按产品线 × 技术类进目录\n标题即可判断主题、型号、作者"),
        ("阅读附件", "PDF/PPT 步骤和配图在页面里\n不必先下载再决定能不能复用"),
        ("作者写作", "源 Sharing 不被覆盖\n日常更新不受整理任务打断"),
        ("知识运营", "增量、去重、限流可续跑\n裂图 17,255 篇已批处理修复"),
    ]
    for i, (t, b) in enumerate(uses):
        x = 360000 + i * 2920000
        _card(s, x, 1100000, 2750000, 4300000)
        _rect(s, x, 1100000, 2750000, 700000, NAVY if i < 3 else ORANGE, radius=0.08)
        _add_text(s, x + 120000, 1250000, 2510000, 450000, t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text(s, x + 180000, 2100000, 2400000, 2800000, b, size=15, color=MUTED)

    # 10 roadmap
    s = nxt("知识生命周期：治理先于问答")
    _add_pic(s, os.path.join(FIG, "fig8_agent_roadmap.png"), 200000, 880000, 11800000, 5450000)

    # 11 close
    s = nxt("可复制，也可继续往下做")
    _add_pic(s, os.path.join(FIG, "fig10_replicate.png"), 200000, 860000, 11800000, 3800000)
    _card(s, 480000, 4800000, 11200000, 1400000)
    _add_text(
        s,
        700000,
        5000000,
        10800000,
        1000000,
        "换 SCAN / TARGET 与标签树，同一套 Agent 可迁到研发、认证、培训知识库。\n"
        "对外发布由协作团队负责；本项目输出已分类、可检索、带质量分的候选。",
        size=16,
        color=NAVY,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    path = os.path.join(OUT_DIR, OUT_NAME)
    prs.save(path)
    return path


if __name__ == "__main__":
    p = build()
    print("saved", p)
    desk = os.path.join(os.path.expanduser("~"), "Desktop")
    dst = os.path.join(desk, OUT_NAME)
    try:
        shutil.copy2(p, dst)
        print("copied", dst)
    except OSError as e:
        print("desktop copy skipped", e)
