# -*- coding: utf-8 -*-
"""One-off: 复审计划书 + 评委问题答复 + PPT。"""

from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu, Inches
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu as PptEmu
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "docs", "AI应用大赛相关材料")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
PPT_NAVY = PptRGB(0x1F, 0x3A, 0x5F)
PPT_ORANGE = PptRGB(0xC4, 0x5C, 0x26)
PPT_BG = PptRGB(0xF7, 0xF5, 0xF0)
PPT_CARD = PptRGB(0xFF, 0xFF, 0xFF)
PPT_MUTED = PptRGB(0x5C, 0x67, 0x70)


def set_run_font(run, *, name="宋体", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    r = run._element.get_or_add_rPr()
    rFonts = r.get_or_add_rFonts()
    latin = "Times New Roman" if name == "宋体" else name
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), name)


def add_p(doc, text, *, size=12, bold=False, center=False, space_after=8, first=True, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.35
    if first:
        pf.first_line_indent = Cm(0.74)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_h(doc, text, *, size=16, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, name="黑体", size=size, bold=True, color=NAVY)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("• " + text)
    set_run_font(run, size=12)
    return p


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): hex_color,
        },
    )
    tcPr.append(shd)


def fill_table(table, rows):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            header = ri == 0
            set_run_font(run, size=10, bold=header)
            if header:
                set_cell_shading(cell, "D9E2F3")


def new_doc():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.8)
    return doc


def write_cover(doc, subtitle: str):
    add_p(doc, "移远通信 AI 实战应用大赛 · 复审材料", size=14, bold=True, center=True, first=False, space_after=4, color=NAVY)
    add_p(doc, subtitle, size=22, bold=True, center=True, first=False, space_after=18, color=NAVY)
    add_p(doc, "项目名称：FAE 知识智能治理平台（AI DocClassifier）", size=12, center=True, first=False, space_after=2)
    add_p(doc, "赛道：运营提效与职能创新　　团队：分类小能手　　部门：技术支持部", size=12, center=True, first=False, space_after=2)
    add_p(doc, "申报日期：2026 年 09 月 17 日　　对应评委意见的补充完善稿", size=12, center=True, first=False, space_after=18)


# ---------------------------------------------------------------------------
# 1) Judge Q&A
# ---------------------------------------------------------------------------

def build_qa() -> str:
    doc = new_doc()
    write_cover(doc, "复审补充：评委问题答复")

    add_p(
        doc,
        "本文专门回应初审评语：「完全依托飞书知识库、缺少原生 Agent、只是做一个 FAE 知识库」。"
        "结论先说清楚：飞书 Wiki 是企业指定的存储与协作层；分类、抽取、标题、编排、质量治理是我们自研的 Agent 流水线。"
        "当前已落地的不是聊天式 RAG，而是让 2 万篇「按人堆放、附件沉睡」的 Sharing 变成可检索、可运营的知识资产——"
        "这是后续 RAG / FAE Support Agent 能做准的前提。",
        first=True,
    )

    add_h(doc, "一、对初审评语的直接回应", size=16)
    add_h(doc, "（一）不是「飞书自带知识库能力」", size=14)
    add_p(
        doc,
        "飞书知识库提供的是：页面编辑、目录、基础全文检索、权限。它不提供、本项目必须自研的能力包括："
        "（1）把 PDF/Word/PPT 附件解析成可检索正文与可显示图片；（2）按 QT-SOP-PM-048E 模组规范约束 LLM，"
        "输出 Product Line × Technical Category 叶子路径；（3）把日期/流水号标题改写成「主题-型号-作者」；"
        "（4）多 Worker 增量扫描、共享去重、飞书/LLM 限流、操作账本；（5）复制后 enrichment（元数据表、图片重绑）。"
        "若只靠飞书自带能力，FAE 仍要在人名文件夹里翻 PDF，标题仍是日期，检索仍找不到附件里的根因。",
    )
    add_h(doc, "（二）原生 Agent 是什么、模型是什么、工具是什么", size=14)
    add_p(
        doc,
        "这里的 Agent 不是对话机器人，而是「感知文档 → 调用工具 → 受约束决策 → 写回知识枢纽」的批处理智能体。"
        "运行时：OpenAI 兼容网关（LLM_BASE_URL，内部 qlitellm）+ 模型 deepseek-v4-flash（可配置 LLM_MODEL）。"
        "每个 Agent 有明确工具集，禁止让模型自由编造分类路径或技术参数。",
    )
    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["Agent", "职责", "模型", "工具（自研/接入）"],
            [
                "抽取 Agent",
                "PDF/Word/PPT → 正文+表+图写入 Wiki 块",
                "规则为主；图需转码不走视觉大模型",
                "PyMuPDF、python-docx、python-pptx；飞书 medias 上传/下载、docx blocks",
            ],
            [
                "分类 Agent",
                "排除过程稿后，输出标签树叶子路径",
                "deepseek-v4-flash",
                "Feishu 扫描/读 raw_content/复制；QT-SOP 正则；LABEL_TREE 约束解码；分类缓存",
            ],
            [
                "标题与元数据 Agent",
                "主题-型号-作者；文内元数据表",
                "deepseek-v4-flash",
                "通讯录/源路径作者回退；Wiki 改标题；docx 表格块",
            ],
            [
                "编排 Agent",
                "增量、去重、限流、账本、控制台派发",
                "不调用 LLM",
                "scan_snapshot、shared_copy_state、tool_ops、Feishu 限速重试",
            ],
            [
                "图片修复 Agent",
                "裂图重绑 / 空图块重提",
                "不调用 LLM",
                "wiki node extra、格式转码、嵌套图块遍历",
            ],
        ],
    )
    add_p(doc, "表 1  已落地的原生 Agent、模型与工具（飞书只提供存储/API，不提供上述决策逻辑）", size=10, first=False, center=True, space_after=12)

    add_h(doc, "二、问题 1：PDF / Word 如何抽取正文、表格、图片", size=16)
    add_p(
        doc,
        "目标不是做通用文档解析 SaaS，而是让附件进入与飞书正文同一条分类和检索链路。"
        "抽取写回源文档「附件：」分区，复制到 TARGET 后再按 wiki 鉴权重绑图片，避免副本裂图。",
    )
    add_h(doc, "（一）PDF（以 PDF 为例的完整路径）", size=14)
    add_bullet(doc, "打开：PyMuPDF（fitz）按页处理。")
    add_bullet(
        doc,
        "正文：page.get_text(\"blocks\")，取文本块（type=0），按块的 y 坐标排序，保持近似阅读顺序；每页结束后插入「第 N 页」分隔。",
    )
    add_bullet(
        doc,
        "表格：当前不单独走 Camelot/版面检测。PDF 表格的单元格会作为空间上相邻的 text block 抽出，写入连续段落。"
        "优点是实现稳定、能进全文检索；局限是不保留栅格表结构。Word 表格则按行抽取为「单元格 | 单元格」文本，结构更好。"
        "评委问的「table」：能检索到表内文字，但 PDF 侧尚未做 HTML/Markdown 表格重建。这是明确的下一阶段优化（见第三节召回）。",
    )
    add_bullet(
        doc,
        "图片：page.get_images 取 XObject；跳过 mask；用页面上的 image rect 决定插入位置（与正文按 y 交织）。"
        "像素经 pixmap 转 RGB JPEG/PNG：CMYK、JPEG2000、超大边（>4096px）、GIF/WEBP/BMP/EMF 一律转码，否则飞书图块会裂图。"
        "上传带 extra.drive_route_token（优先 wiki node token），replace_image 写入宽高。",
    )
    add_h(doc, "（二）Word / PPT", size=14)
    add_bullet(doc, "Word：段落文本；表格按行拼接；DrawingML blip 与 VML v:imagedata（含表格内图片）；.doc 先转 docx。")
    add_bullet(doc, "PPT：遍历形状，含 GROUP 递归；文本框与 PICTURE 按坐标排序后写入。")
    add_h(doc, "（三）是否多模态 RAG", size=14)
    add_p(
        doc,
        "当前不是 CLIP/视觉向量多模态 RAG。处理策略是「多模态摄入、文本层检索」："
        "图进入 Wiki 图块供人阅读与飞书页面展示；表内文字进入正文供全文检索与 LLM 分类。"
        "不做图 embedding、不做页级 OCR 视觉问答。原因：第一瓶颈是 1.7 万篇附件根本进不了索引；"
        "先把图和表变成可展示、可检索的资产，再在治理后的 TARGET 上做向量 RAG 才有意义。"
        "下一阶段：对关键示意图用视觉模型生成 caption 写入正文，再纳入向量索引。",
    )

    add_h(doc, "三、问题 2：抽取之后如何存储、做了哪些优化", size=16)
    add_p(
        doc,
        "存储单元不是向量库，而是「受治理的飞书文档 + 元数据表 + 多维表目录」。这是有意选择：FAE 日常仍在飞书里协作，"
        "治理结果必须落在他们已经打开的入口，而不是另建一套只有开发能用的向量库。",
    )
    add_bullet(doc, "双空间：SCAN 保留作者原文；TARGET 是分类枢纽副本。整理侧改标题/贴表不回写源，避免污染现场写作。")
    add_bullet(doc, "一篇文档的知识包 = 原正文 +「附件提取」分区（文本/表/图）+ 文首元数据表（原名、源路径、作者、产品线、主题、型号、源创建时间）。")
    add_bullet(doc, "对象存储优化：图片绑定 TARGET 的 wiki node，避免 copy 后仍指向 SCAN 媒体 token 导致裂图；存量 17,255 篇已跑修复，重绑约 16.9 万张图，空图块自动从附件重提 698 篇。")
    add_bullet(doc, "噪声过滤：周报/日报/会议纪要/客户问题跟踪不入库，缩小检索宇宙。")
    add_bullet(doc, "去重与增量：shared_copy_state 按 obj_token 全局去重；scan_snapshot 只处理新增；tool_ops 按操作跳过，避免重复贴表/改名。")
    add_bullet(doc, "目录分卷：单层节点过多时按产品线规则 rollover，保证 Wiki 导航可用。")

    add_h(doc, "四、问题 3：召回做了哪些优化", size=16)
    add_p(
        doc,
        "当前召回是「结构化过滤 + 全文检索」，不是 embedding Top-K。对评委关心的 RAG，我们把召回分成已上线和规划中两层，避免把规划说成已上线。",
    )
    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["层级", "已上线（生产）", "作用"],
            ["L0 语料清洗", "排除过程稿；附件正文化；裂图修复", "先让该被搜到的字真的在索引里"],
            ["L1 结构化导航", "Product Line × Technical Category 目录树", "先缩到产品线/技术域，再搜，降低人名目录噪声"],
            ["L2 结果可读", "标题改为主题-型号-作者；元数据表", "搜索命中后 1 秒判断是否值得点开"],
            ["L3 全文", "飞书对 TARGET 正文+抽取区做全文", "PDF 里的 AT 命令、报错码可被搜到"],
            ["L4 目录", "多维表导出元数据/展示标题", "运营侧按作者、型号、路径筛选"],
        ],
    )
    add_p(doc, "表 2  已上线召回优化（尚未上向量 RAG）", size=10, first=False, center=True, space_after=10)
    add_p(
        doc,
        "规划中的 L5：在 TARGET + 质量评分达标语料上建向量索引，做案例问答 / 工单辅助。"
        "分类 Agent 已做的「域强制、叶子约束、Others 纠偏」会直接降低 RAG 的脏召回。"
        "没有 L0–L3，直接上 RAG，模型会检索到周报、裂图 PDF 和日期标题，这正是我们没有一上来做聊天机器人的原因。",
    )

    add_h(doc, "五、除知识库外，还会接入哪些 Agent、用在哪些场景", size=16)
    add_p(
        doc,
        "本项目不是「整个只做一个 RAG 知识库」。它是知识生命周期平台：采集 → 解析 → 分类 →（规划）质量评分 → 人工复核 → 供给下游 Agent。"
        "飞书 Wiki 是底座；上面已经有治理 Agent，下一步才是消费侧 Agent。",
    )
    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["阶段", "Agent / 能力", "场景"],
            ["已上线", "抽取 / 分类 / 标题元数据 / 编排 / 图片修复", "历史 1.7 万篇治理 + 每日增量归档"],
            ["试点中（规则已定义）", "质量评分 Agent（5 维 25 分，必须给证据）", "判断哪些 Sharing 能进对外库或 RAG 白名单"],
            ["下一阶段", "检索问答 Agent（向量 RAG + 目录过滤）", "FAE「这个型号以前怎么排过」"],
            ["下一阶段", "工单辅助 Agent", "打开客户问题时推荐同型号历史案例与步骤"],
            ["后续可选", "Rewrite Agent（不覆盖原文、人工终审）", "B 档文档补步骤/验证，供培训与对外候选"],
        ],
    )
    add_p(doc, "表 3  Agent 路线图：治理必须先于 RAG", size=10, first=False, center=True, space_after=10)
    add_p(
        doc,
        "可复制性：换标签树和 SCAN/TARGET token，同一套 Agent 可迁到研发/认证/培训知识空间。"
        "对外知识库、官网 Qchat、论坛机器人由协作团队建设，本项目输出「已分类、可检索、带质量分」的候选，而不是再造一个飞书。",
    )

    add_h(doc, "六、本轮新增落地数据（复审可核验）", size=16)
    add_bullet(doc, "统一枢纽叶子文档：约 17,105 篇分类复制；标题标准化约 17,092 篇（初审口径）。")
    add_bullet(doc, "2026-09 附件裂图治理：对 TARGET 17,255 篇跑修复；有图 5,823 篇；重绑 169,010 张；空图块自动重提 698 篇；失败 284 张。")
    add_bullet(doc, "提效口径不变：存量约 1 人年已兑现；以后每年约 1.7 FTE，检索复用 1.44 FTE 是大头（公式见计划书）。")

    path = os.path.join(OUT_DIR, "复审补充_评委问题答复_原生Agent与抽取存储召回.docx")
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2) Updated plan (复审版)
# ---------------------------------------------------------------------------

def build_plan() -> str:
    doc = new_doc()
    write_cover(doc, "项目计划书（复审完善版）")

    add_h(doc, "目  录", size=16)
    for line in [
        "一、项目概述",
        "二、项目详细方案（含原生 Agent、抽取、存储、召回）",
        "三、实用性分析",
        "四、完成度与实施计划",
        "五、可复制性与推广价值",
        "六、附件",
    ]:
        add_p(doc, line, first=False, space_after=2)

    add_h(doc, "一、项目概述", size=16)
    add_h(doc, "（一）项目简介", size=14)
    add_p(
        doc,
        "本项目面向全球 FAE Sharing，建设「知识治理 Agent 平台」，而不是再包一层飞书知识库。"
        "系统以飞书 Wiki 为存储入口，自研抽取、分类、标题、编排 Agent："
        "扫描历史与新增文档，解析 PDF/Word/PPT，按 QT-SOP-PM-048E 与标签树完成 Product Line × Technical Category 分类，"
        "写入统一 TARGET 枢纽，并生成可读标题与元数据。约 17,105 篇有效叶子已分类复制，17,092 篇标题标准化；"
        "附件图片已在枢纽侧重绑约 16.9 万张。下一阶段是 5 维 25 分质量评分 Agent，以及在治理语料上的检索问答 Agent。"
        "AI Rewrite 与对外知识库发布由人工确认 / 协作团队负责。",
    )
    add_h(doc, "（二）解决的问题痛点", size=14)
    add_bullet(doc, "人工分类不可持续：历史 Sharing 超 20,000 项；有效叶子约 17,105 篇。按 5 分钟/篇一次性整理约 1,400 人时。")
    add_bullet(doc, "跨产品线检索差：目录按区域/人员沉淀，标题常为日期或流水号。")
    add_bullet(doc, "附件沉睡：有效信息在 PDF/Word/PPT 中，只读飞书正文会漏分类、也搜不到。")
    add_bullet(doc, "飞书不解决治理：原生 Wiki 没有模组约束分类、没有附件结构化、没有增量编排账本。")
    add_bullet(doc, "质量参差：分类完成不等于可复用；缺步骤、验证、适用范围的文档不能直接喂给 RAG。")
    add_h(doc, "（三）应用场景", size=14)
    add_bullet(doc, "历史治理：批量抽取 + 分类 + 标准标题 + 元数据，建成枢纽。")
    add_bullet(doc, "增量治理：只处理新增，形成「写入 → 分类 → 纠偏 → 规则升级」闭环。")
    add_bullet(doc, "检索复用：按产品线/技术类/型号定位历史案例（当前为目录+全文；下一步向量 RAG）。")
    add_bullet(doc, "质量审核（试点）：25 分评分，输出证据与建议。")
    add_bullet(doc, "下游供给：A 档经人工确认后给对外库 / Support Agent，不自动公开。")

    add_h(doc, "二、项目详细方案", size=16)
    add_h(doc, "（一）目标与核心功能", size=14)
    add_p(doc, "总体目标：采集 → 解析 → AI 分类 → 质量评分 → 人工复核 → 高质量候选 → 持续治理。")
    add_bullet(doc, "分类归档 Agent：BFS 叶子文档；规则+LLM 叶子路径；排除周报等过程稿。")
    add_bullet(doc, "多格式抽取 Agent：PDF/Word/PPT 正文、表格文字、图片写入「附件：」区。")
    add_bullet(doc, "标题/元数据 Agent：「主题-型号-作者」；文内 Metadata。")
    add_bullet(doc, "编排 Agent：增量快照、共享去重、API/LLM 限流、控制台。")
    add_bullet(doc, "质量评分 Agent（规则已定义，平台试点）：5 维 25 分，必须给证据。")
    add_h(doc, "（二）原生 Agent、模型与工具（回应「没有 Agent」）", size=14)
    add_p(
        doc,
        "模型：deepseek-v4-flash，经公司 OpenAI 兼容网关调用，温度低、输出 JSON/两行约束格式。"
        "工具：飞书 OpenAPI（wiki 扫描、docx 读写、媒体、复制、多维表）+ PyMuPDF/python-docx/python-pptx +"
        " SQLite 账本。模型不得脱离 LABEL_TREE 和模组正则自由发挥。分层架构：飞书=存储；自研=感知与决策。",
    )
    add_h(doc, "（三）PDF/Word 抽取（回应评委问题 1）", size=14)
    add_p(
        doc,
        "PDF：PyMuPDF 按页 get_text(blocks) 抽正文；表格以空间文本块进入检索（Word 表则按行保留「|」结构）；"
        "图片按 XObject+坐标插入，并转 RGB JPEG/PNG。不是视觉多模态 RAG，而是多模态摄入、文本层检索。"
        "详见同目录《复审补充_评委问题答复》。",
    )
    add_h(doc, "（四）存储与召回优化（回应评委问题 2、3）", size=14)
    add_p(
        doc,
        "存储：SCAN/TARGET 双空间；知识包=正文+抽取区+元数据表；图片按 wiki token 重绑；过程稿不入库。"
        "召回 L0–L4 已上线（清洗、目录树、可读标题、全文、多维表）；L5 向量 RAG 在评分语料就绪后接入。"
        "不把未治理语料直接向量化。",
    )
    add_h(doc, "（五）创新点", size=14)
    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["创新点", "与「只用飞书知识库」的区别", "业务价值"],
            ["受约束分类 Agent", "模组规则先定产品线，LLM 只在子树选叶子", "可运营、可纠偏，不是聊天分类"],
            ["附件结构化摄入", "PDF/PPT 进入正文和图块，而不是当黑盒附件", "分类和检索终于能看见附件知识"],
            ["治理先于 RAG", "先目录+全文+质量分，再向量问答", "避免脏语料污染 Support Agent"],
            ["工程化编排", "增量、去重、限流、账本、控制台", "1.7 万篇可连续跑，而不是 Demo"],
        ],
    )
    add_p(doc, "表 1  创新点对照", size=10, first=False, center=True, space_after=12)

    add_h(doc, "三、实用性分析", size=16)
    add_h(doc, "（一）预期效果（量化）", size=14)
    add_p(
        doc,
        "口径：提效率=(前−后)/前；年等效 FTE=年节省小时/2000。增量 80 篇/周、48 周；检索 150 人×每周 2 次。"
        "存量一次性整理约 1 人年已完成，不计入「每年」。",
        first=True,
    )
    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["对应痛点", "使用前 → 使用后", "提效率 / 年等效 FTE"],
            ["人工分类归档", "5 分钟/篇 → 0.5 分钟/篇", "90%；约 0.14 FTE/年"],
            ["跨人名目录找同类案例", "15 分钟/次 → 3 分钟/次", "80%；约 1.44 FTE/年"],
            ["标题改为主题-型号-作者", "2 分钟/篇 → 0", "100%；约 0.06 FTE/年"],
            ["打开 PDF 才能判断能否复用", "8 分钟/次 → 2 分钟/次", "75%；约 0.04 FTE/年"],
            ["合计（每年持续）", "—", "约 1.7 FTE"],
        ],
    )
    add_p(doc, "表 2  提效（与初审同一公式，未改假设）", size=10, first=False, center=True, space_after=8)
    add_p(doc, "复审新增可核验结果：TARGET 17,255 篇图片修复；重绑 169,010 张；空图块重提 698 篇。", first=True)
    add_h(doc, "（二）用户群体", size=14)
    add_bullet(doc, "一线：全球 FAE。运营：枢纽管理员。间接受益：研发、培训、客户。")
    add_h(doc, "（三）试点", size=14)
    add_p(doc, "分类已生产运行。下一步约 100 篇质量评分试点，20–30 篇黄金样本。RAG Agent 在评分稳定后小流量接入 TARGET。")

    add_h(doc, "四、完成度与实施计划", size=16)
    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    fill_table(
        table,
        [
            ["阶段", "状态", "成果 / 下一步"],
            ["Phase 1 接入验证", "已完成", "Wiki 读取、附件解析、分类链路"],
            ["Phase 2 历史结构化", "核心完成", "约 17,105 篇进 TARGET；17,092 标题；裂图修复"],
            ["Phase 3 持续治理", "运行中", "增量、元数据、Others 纠偏、图片修复"],
            ["Phase 4 质量评分", "规则已定义，试点中", "5 维 25 分；人工复评"],
            ["Phase 5 知识服务 Agent", "规划", "TARGET 上 RAG/工单辅助；对外库由协作团队发布"],
        ],
    )
    add_p(doc, "表 3  完成度", size=10, first=False, center=True, space_after=10)
    add_h(doc, "风险", size=14)
    add_bullet(doc, "API 限流：跨进程限速与重试（含上传 1061045）。")
    add_bullet(doc, "分类偏差：模组规则优先、Others 阈值、重分类搬家。")
    add_bullet(doc, "过早 RAG：未评分语料不进向量库。")
    add_bullet(doc, "密钥：独立飞书 App，.env 不入库。")

    add_h(doc, "五、可复制性与推广价值", size=16)
    add_p(
        doc,
        "可复制的是「接入 → Agent 结构化 → 质量标准 → 人工复核 → 再供给 RAG」流水线，不是某个飞书空间本身。"
        "换源 token 与标签树即可用于研发/认证/培训库。最终目标：每一次客户支持经验，都能经 Agent 理解后被下一位 FAE 复用。",
    )

    add_h(doc, "六、附件", size=16)
    add_p(
        doc,
        "请一并提交：（1）本计划书；（2）《复审补充_评委问题答复》；（3）复审 PPT；"
        "（4）TARGET 目录截图（初审已有）。评分页面与黄金样本报告待试点完成后补。",
        first=True,
    )

    path = os.path.join(OUT_DIR, "复审_项目计划书_FAE知识智能治理平台.docx")
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 3) PPT
# ---------------------------------------------------------------------------

def _set_run(run, *, size=14, bold=False, color=PPT_NAVY, name="Microsoft YaHei"):
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _add_text(slide, l, t, w, h, text, *, size=14, bold=False, color=PPT_NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PptEmu(l), PptEmu(t), PptEmu(w), PptEmu(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return box


def _add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptEmu(l), PptEmu(t), PptEmu(w), PptEmu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def _blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, PPT_BG)
    _add_rect(slide, 0, 0, prs.slide_width, PptEmu(90000), PPT_NAVY)
    return slide


def build_ppt() -> str:
    prs = Presentation()
    prs.slide_width = PptEmu(12192000)
    prs.slide_height = PptEmu(6858000)

    # 1
    s = _blank_slide(prs)
    _add_text(s, 480000, 400000, 11000000, 400000, "FAE 知识智能治理平台  ·  复审", size=14, bold=True, color=PptRGB(255, 255, 255))
    _add_text(s, 480000, 1400000, 11000000, 900000, "不是飞书知识库套壳\n是已上线的文档治理 Agent 流水线", size=32, bold=True)
    _add_text(
        s,
        480000,
        3400000,
        11000000,
        800000,
        "评委关注：原生 Agent / 模型与工具 / PDF 抽取 / 存储 / 召回 / 是否只是 RAG",
        size=16,
        color=PPT_MUTED,
    )
    _add_text(
        s,
        480000,
        5000000,
        11000000,
        700000,
        "17,105 篇已分类  ·  17,092 标题标准化  ·  169,010 张附件图重绑  ·  deepseek-v4-flash",
        size=14,
        bold=True,
        color=PPT_ORANGE,
    )

    # 2
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "评委问题 → 一句话回答", size=14, bold=True, color=PptRGB(255, 255, 255))
    cards = [
        (480000, "飞书之外还做什么", "自研抽取、分类、标题、编排 Agent。Wiki 只负责存和协同。"),
        (3300000, "模型 / 工具", "deepseek-v4-flash + 飞书 OpenAPI + PyMuPDF/docx/pptx + 账本。"),
        (6120000, "是不是只有 RAG", "现在是治理平台。RAG/Support 是下一阶段，且必须吃治理后语料。"),
        (8940000, "为何算创新", "2 万篇按人堆放+附件沉睡，飞书原生存不了；约束分类+抽取是自研。"),
    ]
    for x, title, body in cards:
        _add_rect(s, x, 1200000, 2700000, 4200000, PPT_CARD)
        _add_rect(s, x, 1200000, 2700000, 120000, PPT_ORANGE)
        _add_text(s, x + 140000, 1500000, 2400000, 900000, title, size=16, bold=True)
        _add_text(s, x + 140000, 2500000, 2400000, 2500000, body, size=13, color=PPT_MUTED)

    # 3
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "原生 Agent 栈（已上线）", size=14, bold=True, color=PptRGB(255, 255, 255))
    rows = [
        ("抽取 Agent", "PDF/Word/PPT → 正文+表文字+图"),
        ("分类 Agent", "QT-SOP 正则 + 标签树约束 LLM"),
        ("标题/元数据 Agent", "主题-型号-作者 + 文内表"),
        ("编排 Agent", "增量 · 去重 · 限流 · 控制台"),
        ("修复 Agent", "裂图重绑 / 空图块重提"),
    ]
    for i, (a, b) in enumerate(rows):
        y = 1100000 + i * 950000
        _add_rect(s, 480000, y, 11200000, 850000, PPT_CARD)
        _add_text(s, 700000, y + 180000, 3500000, 500000, a, size=18, bold=True, color=PPT_ORANGE)
        _add_text(s, 4300000, y + 180000, 7000000, 500000, b, size=16)

    # 4 PDF
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "问题1  PDF 怎么抽正文 / 表 / 图", size=14, bold=True, color=PptRGB(255, 255, 255))
    items = [
        ("正文", "PyMuPDF get_text(blocks)，按 y 排序，分页写入 Wiki 段落。"),
        ("表格", "PDF：单元格作为文本块进入检索（不重建栅格）。Word：按行「A | B | C」。"),
        ("图片", "XObject → RGB JPEG/PNG → 飞书图块；wiki token 鉴权，避免副本裂图。"),
        ("多模态 RAG？", "否。当前是多模态摄入 + 文本检索。Caption/向量是下一步。"),
    ]
    for i, (a, b) in enumerate(items):
        y = 1100000 + i * 1200000
        _add_rect(s, 480000, y, 11200000, 1050000, PPT_CARD)
        _add_text(s, 700000, y + 250000, 2800000, 500000, a, size=18, bold=True)
        _add_text(s, 3700000, y + 250000, 7600000, 600000, b, size=16, color=PPT_MUTED)

    # 5 storage + recall
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "问题2–3  存储与召回", size=14, bold=True, color=PptRGB(255, 255, 255))
    _add_rect(s, 480000, 1100000, 5400000, 5000000, PPT_CARD)
    _add_text(s, 700000, 1300000, 5000000, 400000, "存储优化", size=18, bold=True, color=PPT_ORANGE)
    for i, t in enumerate(
        [
            "SCAN 写作 / TARGET 治理，互不覆盖",
            "一篇 = 正文 + 抽取区 + 元数据表",
            "图片按 TARGET wiki 重绑（16.9 万张）",
            "过程稿不入库；obj_token 去重",
            "不是先建向量库",
        ]
    ):
        _add_text(s, 700000, 1900000 + i * 700000, 4900000, 600000, "·  " + t, size=15)
    _add_rect(s, 6200000, 1100000, 5500000, 5000000, PPT_CARD)
    _add_text(s, 6400000, 1300000, 5100000, 400000, "召回优化（已上线 L0–L4）", size=18, bold=True, color=PPT_ORANGE)
    for i, t in enumerate(
        [
            "L0 附件正文化、裂图修复",
            "L1 产品线 × 技术类目录过滤",
            "L2 标题可读（主题-型号-作者）",
            "L3 飞书全文（含 PDF 抽出的字）",
            "L5 向量 RAG = 规划，吃评分语料",
        ]
    ):
        _add_text(s, 6400000, 1900000 + i * 700000, 5100000, 600000, "·  " + t, size=15)

    # 6 roadmap
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "不是只有 RAG：Agent 场景路线", size=14, bold=True, color=PptRGB(255, 255, 255))
    phases = [
        ("已上线", "治理 Agent\n抽取·分类·标题·编排"),
        ("试点", "质量评分 Agent\n5 维 25 分+证据"),
        ("下一步", "检索问答 Agent\nTARGET 向量 RAG"),
        ("下一步", "工单辅助 Agent\n同型号历史案例"),
        ("可选", "Rewrite Agent\n人工终审不覆盖原文"),
    ]
    for i, (h, b) in enumerate(phases):
        x = 350000 + i * 2300000
        _add_rect(s, x, 1400000, 2100000, 3800000, PPT_CARD)
        _add_rect(s, x, 1400000, 2100000, 120000, PPT_ORANGE if i < 2 else PPT_NAVY)
        _add_text(s, x + 120000, 1700000, 1860000, 600000, h, size=16, bold=True, color=PPT_ORANGE)
        _add_text(s, x + 120000, 2500000, 1860000, 2200000, b, size=14, color=PPT_MUTED)
    _add_text(s, 480000, 5600000, 11000000, 500000, "没有 L0–L3 治理，RAG 会检索到周报、裂图和日期标题。所以先治理，再问答。", size=14, color=PPT_MUTED)

    # 7 numbers
    s = _blank_slide(prs)
    _add_text(s, 480000, 250000, 11000000, 350000, "可核验落地 + 提效口径", size=14, bold=True, color=PptRGB(255, 255, 255))
    nums = [
        ("17,105", "分类进枢纽"),
        ("17,092", "标准标题"),
        ("169,010", "图重绑"),
        ("~1.7", "FTE / 年"),
    ]
    for i, (n, lab) in enumerate(nums):
        x = 480000 + i * 2900000
        _add_rect(s, x, 1200000, 2700000, 2200000, PPT_CARD)
        _add_text(s, x + 100000, 1450000, 2500000, 900000, n, size=32, bold=True, color=PPT_ORANGE, align=PP_ALIGN.CENTER)
        _add_text(s, x + 100000, 2500000, 2500000, 500000, lab, size=14, align=PP_ALIGN.CENTER, color=PPT_MUTED)
    _add_text(
        s,
        480000,
        3800000,
        11200000,
        2000000,
        "归档 5→0.5 分钟  ·  检索 15→3 分钟  ·  标题 2→0  ·  打开附件 8→2\n"
        "存量约 1 人年已完成；每年约 1.7 FTE，检索 1.44 是大头。FTE = 节省小时 / 2000。",
        size=16,
        color=PPT_NAVY,
    )

    path = os.path.join(OUT_DIR, "复审_PPT_FAE知识智能治理平台.pptx")
    prs.save(path)
    return path


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    qa = build_qa()
    plan = build_plan()
    ppt = build_ppt()
    print(qa)
    print(plan)
    print(ppt)


if __name__ == "__main__":
    main()
